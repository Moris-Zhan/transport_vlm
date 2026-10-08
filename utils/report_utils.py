"""Vision-LLM 推論結果的 report 產出工具。

從 vlm_infer_test.ipynb 搬移而來:讀取標註 jsonl 資料,並將推論結果連同
準確率統計(呼叫 utils.accuracy_utils)彙整成含圖片的 Excel report。
"""

import json
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage

from utils.accuracy_utils import compare_texts_by_characters, extract_pointer, compare_extraction


def read_json_lines(file_path):
    """讀取包含多行 JSON 的檔案"""
    data_list = []
    with open(file_path, 'r', encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if line:
                data_list.append(json.loads(line))
    return data_list


def parse_json_strings(json_strings):
    """從 JSON 字串列表解析資料"""
    return [json.loads(s) for s in json_strings]


def create_excel_with_images_inf_data(data_list, data_sources, llm_outputs, output_file='output.xlsx'):
    """
    建立包含圖片的 Excel 檔案

    參數:
        data_list: JSON 資料列表
        output_file: 輸出的 Excel 檔案名稱
    """
    # 建立工作簿和工作表
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inference Data"
    total_accuracy, qianmo_accuracy, transport_accuracy = [], [], []
    square_accuracy, skew_accuracy = [], []
    ocr_accuracy = []
    ocr_precision, ocr_recall, ocr_f1 = [], [], []

    # 設定標題列
    headers = ['Background', 'Image Path', 'Source', 'Text', 'Inference', 'Picture',
               'Text_Accuracy(%)', 'Icons_Accuracy(%)', 'OCR_Accuracy(%)',
               'OCR_Precision(%)', 'OCR_Recall(%)', 'OCR_F1(%)',
               'Compare', 'OCR_Report']
    ws.append(headers)

    # 設定標題列樣式
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = openpyxl.styles.Font(bold=True)

    # 設定欄寬
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 37
    ws.column_dimensions['C'].width = 8
    ws.column_dimensions['D'].width = 80
    ws.column_dimensions['E'].width = 80
    ws.column_dimensions['F'].width = 80

    ws.column_dimensions['G'].width = 15
    ws.column_dimensions['H'].width = 15
    ws.column_dimensions['I'].width = 15
    ws.column_dimensions['J'].width = 15
    ws.column_dimensions['K'].width = 15
    ws.column_dimensions['L'].width = 15
    ws.column_dimensions['M'].width = 80
    ws.column_dimensions['N'].width = 80

    # 填入資料並插入圖片
    for idx, data in enumerate(data_list, start=2):
        # 填入文字資料
        background_cell = ws.cell(row=idx, column=1, value=data.get('background', ''))
        background_cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')
        image_cell = ws.cell(row=idx, column=2, value=data.get('image', ''))
        image_cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')

        # 填入 legend_source 欄位並設定自動換行
        legend_source = data_sources[idx-2]
        source_cell = ws.cell(row=idx, column=3, value=legend_source)
        source_cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')

        # 填入 Text 欄位並設定自動換行
        text_cell = ws.cell(row=idx, column=4, value=data.get('text', ''))
        text_cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')

        # 填入 llm_output 欄位並設定自動換行
        llm_output = llm_outputs[idx-2]
        llm_cell = ws.cell(row=idx, column=5, value=llm_output)
        llm_cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')

        llm_answer = data.get('text', '')
        report_line, accuracy = compare_texts_by_characters(llm_answer, llm_output)

        # 填入 Text_Accuracy(%) 欄位並設定自動換行
        accuracy_cell = ws.cell(row=idx, column=7, value=accuracy)
        accuracy_cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')
        # ---------------------------------------------------------------------------------------
        # 細指標比對
        icons_lbl = data.get('icons', '')
        ocr_nums_lbl = data.get('ocr_nums', '')
        ocr_labels_lbl = data.get('ocr_labels', '')
        lbl_result = {
            "icons_lbl": icons_lbl,
            "ocr_nums_lbl": ocr_nums_lbl,
            "ocr_labels_lbl": ocr_labels_lbl,
        }

        (icons, ocr_nums, ocr_labels) = extract_pointer(llm_output)
        # 結果輸出
        pred = {
            "icons": icons,
            "ocr_nums": ocr_nums,
            "ocr_labels": ocr_labels,
        }
        # 1️⃣ icon 準確率
        result, ocr_report = compare_extraction(lbl_result, pred)

        # 填入 Icons_Accuracy(%) 欄位並設定自動換行
        icons_accuracy_cell = ws.cell(row=idx, column=8, value=result["icons_accuracy"])
        icons_accuracy_cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')

        # 填入 OCR_Accuracy(%) 欄位並設定自動換行
        ocr_accuracy_cell = ws.cell(row=idx, column=9, value=result["ocr_accuracy"])
        ocr_accuracy_cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')

        # 填入 OCR_Precision(%)/OCR_Recall(%)/OCR_F1(%) 欄位並設定自動換行
        ocr_precision_cell = ws.cell(row=idx, column=10, value=result["ocr_precision"])
        ocr_precision_cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')
        ocr_recall_cell = ws.cell(row=idx, column=11, value=result["ocr_recall"])
        ocr_recall_cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')
        ocr_f1_cell = ws.cell(row=idx, column=12, value=result["f1"])
        ocr_f1_cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')
        # ---------------------------------------------------------------------------------------
        # 填入 comp_cell 欄位並設定自動換行
        comp_cell = ws.cell(row=idx, column=13, value=report_line)
        comp_cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')

        ocr_report_cell = ws.cell(row=idx, column=14, value=ocr_report)
        ocr_report_cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')

        total_accuracy.append(accuracy)
        ocr_accuracy.append(result["ocr_accuracy"])
        ocr_precision.append(result["ocr_precision"])
        ocr_recall.append(result["ocr_recall"])
        ocr_f1.append(result["f1"])
        if legend_source == "阡陌":
            qianmo_accuracy.append(accuracy)
        elif legend_source == "運研所":
            transport_accuracy.append(accuracy)

        road_type_lbl = data.get('road_type', '')
        if road_type_lbl == "正交":
            square_accuracy.append(accuracy)
        elif road_type_lbl == "斜交":
            skew_accuracy.append(accuracy)

        # 插入圖片
        image_path = data.get('image', '')
        if image_path and Path(image_path).exists():
            try:
                # 載入圖片
                img = OpenpyxlImage(image_path)

                # 調整圖片大小為 500x500
                img.width = 500
                img.height = 500

                # 設定列高以容納圖片
                ws.row_dimensions[idx].height = img.height * 0.75

                # 將圖片錨定到 F 欄的儲存格
                cell_position = f'F{idx}'
                img.anchor = cell_position

                # 添加圖片到工作表
                ws.add_image(img)
            except Exception as e:
                ws.cell(row=idx, column=6, value=f'圖片載入失敗: {str(e)}')
                print(f"警告: 無法載入圖片 {image_path}: {str(e)}")
        else:
            ws.cell(row=idx, column=6, value='圖片不存在')
            print(f"警告: 找不到圖片檔案 {image_path}")

    # 在最後一列新增總體準確率
    if total_accuracy:
        avg_accuracy = sum(total_accuracy) / len(total_accuracy)
        if len(qianmo_accuracy) > 0:
            qianmo_accuracy = sum(qianmo_accuracy) / len(qianmo_accuracy)
        transport_accuracy = sum(transport_accuracy) / len(transport_accuracy)
        square_accuracy = sum(square_accuracy) / len(square_accuracy)
        skew_accuracy = sum(skew_accuracy) / len(skew_accuracy)
        ocr_accuracy = sum(ocr_accuracy) / len(ocr_accuracy)
        ocr_precision = sum(ocr_precision) / len(ocr_precision)
        ocr_recall = sum(ocr_recall) / len(ocr_recall)
        ocr_f1 = sum(ocr_f1) / len(ocr_f1)

        total_row = len(data_list) + 2
        ws.cell(row=total_row, column=7, value=f'Total accuracy: {avg_accuracy:.2f}%')
        if len(qianmo_accuracy) > 0:
            ws.cell(row=total_row+1, column=7, value=f'阡陌 accuracy: {qianmo_accuracy:.2f}%')
        ws.cell(row=total_row+2, column=7, value=f'運研所 accuracy: {transport_accuracy:.2f}%')
        ws.cell(row=total_row+3, column=7, value=f'正交 accuracy: {square_accuracy:.2f}%')
        ws.cell(row=total_row+4, column=7, value=f'斜交 accuracy: {skew_accuracy:.2f}%')
        ws.cell(row=total_row+5, column=7, value=f'OCR accuracy: {ocr_accuracy:.2f}%')
        ws.cell(row=total_row+6, column=7, value=f'OCR precision: {ocr_precision:.2f}%')
        ws.cell(row=total_row+7, column=7, value=f'OCR recall: {ocr_recall:.2f}%')
        ws.cell(row=total_row+8, column=7, value=f'OCR F1: {ocr_f1:.2f}%')
        for i in range(9):
            ws.cell(row=total_row + i, column=7).font = openpyxl.styles.Font(bold=True)

    # 儲存檔案
    wb.save(output_file)
    print(f"Excel 檔案已成功建立: {output_file}")
    print(f"Total accuracy: {avg_accuracy :.2f}%")
    if len(qianmo_accuracy) > 0:
        print(f"阡陌 accuracy: {qianmo_accuracy :.2f}%")
    print(f"運研所 accuracy: {transport_accuracy :.2f}%")
    print(f"正交 accuracy: {square_accuracy :.2f}%")
    print(f"斜交 accuracy: {skew_accuracy :.2f}%")
    print(f"OCR accuracy: {ocr_accuracy :.2f}%")
    print(f"OCR precision: {ocr_precision :.2f}%")
    print(f"OCR recall: {ocr_recall :.2f}%")
    print(f"OCR F1: {ocr_f1 :.2f}%")
